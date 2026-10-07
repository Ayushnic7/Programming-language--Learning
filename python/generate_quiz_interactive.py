import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_python_quiz_ppt():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    BG_DARK = RGBColor(15, 23, 42)      # #0F172A
    ACCENT_BLUE = RGBColor(14, 165, 233)  # #0EA5E9
    TEXT_LIGHT = RGBColor(248, 250, 252) # #F8FAFC
    CARD_BG = RGBColor(30, 41, 59)      # #1E293B
    ANSWER_BG = RGBColor(22, 101, 52)   # #166534

    questions = [
        # Chapter 1: Variables & Operators
        ("In Python, which key combination is used to comment out code in VS Code?", ["A) Shift + /", "B) Ctrl + /", "C) Alt + /", "D) Tab + /"], "B) Ctrl + /"),
        ("Which bracket symbol is referred to as 'Parenthesis' in Python?", ["A) [ ]", "B) { }", "C) ( )", "D) < >"], "C) ( )"),
        ("What symbol is used to define a block of code (functions, loops, if statements)?", ["A) Semicolon (;)", "B) Colon (:)", "C) Comma (,)", "D) Hash (#)"], "B) Colon (:)"),
        ("What is the output of `print(a % b)` when `a = 4` and `b = 2`?", ["A) 2", "B) 0", "C) 8", "D) 16"], "B) 0"),
        ("What does the exponentiation operator `**` do in `5 ** 2`?", ["A) Multiplies 5 by 2", "B) Divides 5 by 2", "C) Raises 5 to the power 2", "D) Finds remainder of 5/2"], "C) Raises 5 to the power 2"),
        ("What is the result of `print(5 != 2)` in Python?", ["A) False", "B) True", "C) None", "D) SyntaxError"], "B) True"),
        ("What does the relational operator `<=` mean?", ["A) Less than", "B) Greater than or equal to", "C) Less than or equal to", "D) Not equal to"], "C) Less than or equal to"),
        ("If `num = 10`, what is the value of `num` after `num += 10`?", ["A) 10", "B) 0", "C) 20", "D) 100"], "C) 20"),
        ("What does `print(not False)` evaluate to in Python?", ["A) False", "B) True", "C) None", "D) Error"], "B) True"),
        ("Which logical operator returns `True` only if BOTH conditions are true?", ["A) or", "B) not", "C) and", "D) in"], "C) and"),
        ("Which logical operator returns `True` if AT LEAST ONE condition is true?", ["A) and", "B) not", "C) or", "D) is"], "C) or"),
        ("When Python automatically changes data type during operations, it is called:", ["A) Type Casting", "B) Type Conversion", "C) Type Checking", "D) Type Decorating"], "B) Type Conversion"),
        ("Explicitly converting a variable's data type manually is called:", ["A) Automatic Conversion", "B) Type Casting", "C) Dynamic Typing", "D) Overloading"], "B) Type Casting"),
        ("What data type does the `input()` function always return by default?", ["A) int", "B) float", "C) str", "D) bool"], "C) str"),
        ("How do you properly take an integer input from a user in Python?", ["A) input(int(\"...\"))", "B) int(input(\"...\"))", "C) integer(\"...\")", "D) str(input(\"...\"))"], "B) int(input(\"...\"))"),
        ("What is the output of `type(3.14)`?", ["A) <class 'int'>", "B) <class 'str'>", "C) <class 'float'>", "D) <class 'double'>"], "C) <class 'float'>"),
        ("How is area of a square calculated if `side = 5` using the power operator?", ["A) side * 2", "B) side ** 2", "C) side % 2", "D) side + side"], "B) side ** 2"),
        ("What is the value of `(5 + 5) / 2` in Python?", ["A) 5", "B) 5.0", "C) 10", "D) 2.5"], "B) 5.0"),
        ("What result does `print(10 >= 10)` return?", ["A) False", "B) True", "C) 10", "D) Error"], "B) True"),
        ("Which of the following is a valid Python variable name?", ["A) 1name", "B) my_name", "C) my-name", "D) class"], "B) my_name"),

        # Chapter 2: Strings & Conditionals
        ("What process is used to join two strings together using `+`?", ["A) Slicing", "B) Indexing", "C) Concatenation", "D) Casting"], "C) Concatenation"),
        ("How are spaces treated in `len(\"hello world\")`?", ["A) Ignored", "B) Counted as 1 character", "C) Raises Error", "D) Converts to zero"], "B) Counted as 1 character"),
        ("What is the starting index of a string in Python?", ["A) 1", "B) -1", "C) 0", "D) Any number"], "C) 0"),
        ("If `str = \"asunix\"`, what is `str[0]`?", ["A) 'a'", "B) 's'", "C) 'x'", "D) 'u'"], "A) 'a'"),
        ("In slicing `str[start:end]`, the `end` index is:", ["A) Inclusive", "B) Exclusive", "C) Mandatory", "D) Ignored"], "B) Exclusive"),
        ("What does `\"apna college\"[0:4]` return?", ["A) \"apna\"", "B) \"apn\"", "C) \"apna \"", "D) \"college\""], "A) \"apna\""),
        ("What index is used to access the very last character of a string?", ["A) 0", "B) len(str)", "C) -1", "D) -0"], "C) -1"),
        ("Which string method converts all characters in a string to uppercase?", ["A) upper()", "B) capitalize()", "C) top()", "D) uppercase()"], "A) upper()"),
        ("What does `str.strip()` do?", ["A) Deletes string", "B) Removes leading & trailing spaces", "C) Splits string into list", "D) Converts to lower"], "B) Removes leading & trailing spaces"),
        ("Which method replaces occurrences of a substring with a new substring?", ["A) change()", "B) replace()", "C) swap()", "D) set()"], "B) replace()"),
        ("What does `str.count(\"$\")` do?", ["A) Checks if string starts with $", "B) Returns index of $", "C) Counts occurrences of $", "D) Replaces $"], "C) Counts occurrences of $"),
        ("Which method capitalizes ONLY the first character of a string?", ["A) upper()", "B) title()", "C) capitalize()", "D) firstUpper()"], "C) capitalize()"),
        ("Which conditional statement keyword is checked ONLY IF the previous `if` is False?", ["A) else", "B) elif", "C) then", "D) condition"], "B) elif"),
        ("What happens if an `if` condition is `True` in an `if-elif-else` chain?", ["A) Executes all elif blocks", "B) Executes block & skips remaining elif/else", "C) Executes else block", "D) Restarts program"], "B) Executes block & skips remaining elif/else"),
        ("Writing an `if` statement inside another `if` statement is called:", ["A) Chaining", "B) Nesting", "C) Looping", "D) Joining"], "B) Nesting"),
        ("What is the output for `age = 80` in:\n`if age >= 18: if age >= 80: print('cannot drive') else: print('can drive')`?", ["A) can drive", "B) cannot drive", "C) Error", "D) No output"], "B) cannot drive"),
        ("How do you check if a number `num` is even in Python?", ["A) num / 2 == 0", "B) num % 2 == 0", "C) num % 2 == 1", "D) num ** 2 == 0"], "B) num % 2 == 0"),
        ("What operator is used to check if a number is a multiple of 7?", ["A) num / 7 == 0", "B) num % 7 == 0", "C) num * 7 == 0", "D) num // 7 == 0"], "B) num % 7 == 0"),
        ("What does `str.endswith(\"salt\")` return?", ["A) Index position", "B) String slice", "C) Boolean (True/False)", "D) Error"], "C) Boolean (True/False)"),
        ("If `marks = 85`, which condition executes in grade check (`>=90`:A, `>=80`:B, `>=70`:C)?", ["A) Grade A", "B) Grade B", "C) Grade C", "D) Grade D"], "B) Grade B"),

        # Chapter 3: Lists & Tuples
        ("Which bracket is used to define a List in Python?", ["A) Parentheses ()", "B) Square brackets []", "C) Curly braces {}", "D) Angle brackets <>"], "B) Square brackets []"),
        ("Lists in Python are mutable. What does 'mutable' mean?", ["A) Values cannot be changed", "B) Values can be changed", "C) Can only store numbers", "D) Fixed size"], "B) Values can be changed"),
        ("Strings are immutable, whereas Lists are:", ["A) Immutable", "B) Mutable", "C) Constant", "D) Unordered"], "B) Mutable"),
        ("What does `list.append(6)` do?", ["A) Adds 6 at index 0", "B) Adds 6 at the end of the list", "C) Deletes element 6", "D) Replaces index 6"], "B) Adds 6 at the end of the list"),
        ("How do you sort a list in descending order?", ["A) list.sort(reverse=True)", "B) list.sort(descending=True)", "C) list.reverse_sort()", "D) list.sort(-1)"], "A) list.sort(reverse=True)"),
        ("What does `list.reverse()` do?", ["A) Sorts list", "B) Clears list", "C) Reverses elements in-place", "D) Creates new empty list"], "C) Reverses elements in-place"),
        ("In `list.insert(3, 9)`, what does `3` represent?", ["A) Value to insert", "B) Index position to insert at", "C) Number of times to repeat", "D) Element to delete"], "B) Index position to insert at"),
        ("What does `list.remove(4)` do?", ["A) Removes element at index 4", "B) Removes first occurrence of value 4", "C) Removes last 4 items", "D) Raises syntax error"], "B) Removes first occurrence of value 4"),
        ("What does `list.pop(4)` do?", ["A) Removes item at index 4", "B) Removes value 4", "C) Adds 4 to end", "D) Reverses 4 items"], "A) Removes item at index 4"),
        ("Which symbol defines a Tuple?", ["A) []", "B) {}", "C) ()", "D) <>"], "C) ()"),
        ("Tuples in Python are:", ["A) Mutable", "B) Immutable", "C) Unordered", "D) Dynamic"], "B) Immutable"),
        ("What is the data type of `tup = (1)`?", ["A) tuple", "B) int", "C) list", "D) set"], "B) int"),
        ("How do you create a single-element tuple correctly?", ["A) tup = (1)", "B) tup = (1,)", "C) tup = [1]", "D) tup = {1}"], "B) tup = (1,)"),
        ("What does `tup.index(3)` return?", ["A) Count of 3", "B) Index of first occurrence of 3", "C) True/False", "D) Value after 3"], "B) Index of first occurrence of 3"),
        ("What does `tup.count(3)` return?", ["A) Index of 3", "B) Number of occurrences of 3", "C) Total size of tuple", "D) Error if 3 missing"], "B) Number of occurrences of 3"),
        ("What is a Palindrome?", ["A) Reads same forward & backward", "B) Number divisible by 2", "C) List with duplicate keys", "D) Sorted sequence"], "A) Reads same forward & backward"),
        ("How can you copy a list `l1` in Python?", ["A) l2 = l1.copy()", "B) l2 = l1.clone()", "C) l2 = copy(l1)", "D) l2 = l1.duplicate()"], "A) l2 = l1.copy()"),
        ("What is the result of `[\"a\", \"b\"] + [\"c\"]`?", ["A) [\"a\", \"b\", \"c\"]", "B) [\"abc\"]", "C) Error", "D) [[\"a\", \"b\"], [\"c\"]]"], "A) [\"a\", \"b\", \"c\"]"),
        ("What happens if you try to change `tup[0] = 5`?", ["A) Changes value", "B) TypeError", "C) Appends 5", "D) Converts to list"], "B) TypeError"),
        ("What does `len(tup)` return?", ["A) Maximum element", "B) Total number of elements", "C) Memory size", "D) Last index"], "B) Total number of elements"),

        # Chapter 4: Dictionaries & Sets
        ("Dictionaries store data in which format?", ["A) Index-Value pairs", "B) Key-Value pairs", "C) Ordered Lists", "D) Single Elements"], "B) Key-Value pairs"),
        ("How do you access the value associated with key 'name' in dictionary `info`?", ["A) info('name')", "B) info[\"name\"]", "C) info.name", "D) info->name"], "B) info[\"name\"]"),
        ("Are dictionary keys allowed to be duplicated?", ["A) Yes", "B) No", "C) Only integer keys", "D) Only string keys"], "B) No"),
        ("Which brackets are used to define a Dictionary?", ["A) []", "B) ()", "C) {}", "D) <>"], "C) {}"),
        ("How do you create an empty dictionary?", ["A) dict = []", "B) dict = ()", "C) dict = {}", "D) dict = set()"], "C) dict = {}"),
        ("What does `dict.keys()` return?", ["A) All values", "B) All key-value pairs", "C) All keys", "D) Length of dictionary"], "C) All keys"),
        ("What does `dict.values()` return?", ["A) All keys", "B) All values", "C) First key only", "D) Dictionary type"], "B) All values"),
        ("What does `dict.items()` return?", ["A) List of tuples of key-value pairs", "B) Keys only", "C) Values only", "D) Total count"], "A) List of tuples of key-value pairs"),
        ("What happens when you use `dict.get(\"key\")` for a key that DOES NOT exist?", ["A) Raises KeyError", "B) Returns None", "C) Returns False", "D) Crashes program"], "B) Returns None"),
        ("What happens when you use `dict[\"key\"]` for a key that DOES NOT exist?", ["A) Returns None", "B) Raises KeyError", "C) Creates key with value 0", "D) Returns empty string"], "B) Raises KeyError"),
        ("How do you add or update a key-value pair in a dictionary?", ["A) dict.add(k, v)", "B) dict[k] = v", "C) dict.insert(k, v)", "D) dict.append(k, v)"], "B) dict[k] = v"),
        ("A dictionary inside another dictionary is called:", ["A) Double Dict", "B) Nested Dictionary", "C) Linked Dictionary", "D) Extended Dict"], "B) Nested Dictionary"),
        ("What is a Set in Python?", ["A) Ordered collection of duplicates", "B) Unordered collection of unique elements", "C) Key-value mapping", "D) Immutable list"], "B) Unordered collection of unique elements"),
        ("How do you create an empty set in Python?", ["A) s = {}", "B) s = set()", "C) s = []", "D) s = ()"], "B) s = set()"),
        ("What happens when duplicate elements are added to a Set?", ["A) Throws Error", "B) Duplicates are ignored", "C) Overwrites set", "D) Stored as list"], "B) Duplicates are ignored"),
        ("Which method adds an element to a Set?", ["A) set.append()", "B) set.add()", "C) set.insert()", "D) set.push()"], "B) set.add()"),
        ("Which method combines elements from two sets returning ONLY unique elements?", ["A) intersection()", "B) union()", "C) difference()", "D) concat()"], "B) union()"),
        ("Which method returns common elements present in BOTH sets?", ["A) union()", "B) intersection()", "C) difference()", "D) symmetric_difference()"], "B) intersection()"),
        ("What does `set.clear()` do?", ["A) Deletes set variable", "B) Empties all elements from set", "C) Removes last element", "D) Resets elements to 0"], "B) Empties all elements from set"),
        ("Are Set elements mutable or indexable via `set[0]`?", ["A) Yes, fully indexable", "B) No, sets are unordered & not indexable", "C) Only integer sets", "D) Yes, with reverse indexing"], "B) No, sets are unordered & not indexable"),

        # Extra Questions to complete 100 MCQs
        ("What is the output of `bool(\"\")` in Python?", ["A) True", "B) False", "C) None", "D) Error"], "B) False"),
        ("What is the output of `bool(\"hello\")`?", ["A) False", "B) True", "C) 1", "D) None"], "B) True"),
        ("Which of the following data types is IMMUTABLE?", ["A) List", "B) Dictionary", "C) Tuple", "D) Set"], "C) Tuple"),
        ("Which of the following data types is MUTABLE?", ["A) String", "B) Tuple", "C) Int", "D) List"], "D) List"),
        ("What result does `3 * \"a\"` produce in Python?", ["A) Error", "B) \"aaa\"", "C) \"a3\"", "D) 3"], "B) \"aaa\""),
        ("What is integer division operator in Python?", ["A) /", "B) //", "C) %", "D) div"], "B) //"),
        ("What is the output of `7 // 2`?", ["A) 3.5", "B) 3", "C) 4", "D) 1"], "B) 3"),
        ("What is the output of `7 % 2`?", ["A) 3.5", "B) 3", "C) 1", "D) 0"], "C) 1"),
        ("How do you find the data type of any object in Python?", ["A) typeof(obj)", "B) type(obj)", "C) datatype(obj)", "D) class(obj)"], "B) type(obj)"),
        ("What is the default value returned by a function that doesn't return anything?", ["A) 0", "B) False", "C) None", "D) Empty string"], "C) None"),
        ("Which keyword is used to start a function definition in Python?", ["A) function", "B) def", "C) define", "D) fun"], "B) def"),
        ("What does the `len()` function return for a dictionary `{\"a\": 1, \"b\": 2}`?", ["A) 1", "B) 2", "C) 4", "D) Error"], "B) 2"),
        ("Which statement is used to exit a loop prematurely?", ["A) stop", "B) break", "C) exit", "D) return"], "B) break"),
        ("Which statement skips the rest of current loop iteration and moves to next?", ["A) skip", "B) pass", "C) continue", "D) next"], "C) continue"),
        ("Which placeholder statement does nothing in Python syntax?", ["A) pass", "B) continue", "C) null", "D) break"], "A) pass"),
        ("What does `str.lower()` return if string is already lowercase?", ["A) New string unmodified", "B) Error", "C) None", "D) Upper string"], "A) New string unmodified"),
        ("What is the result of `10 == \"10\"` in Python?", ["A) True", "B) False", "C) TypeError", "D) None"], "B) False"),
        ("Which operator tests object identity (same memory location)?", ["A) ==", "B) is", "C) in", "D) equals"], "B) is"),
        ("Which operator checks membership of an item in a sequence?", ["A) contains", "B) in", "C) has", "D) exists"], "B) in"),
        ("What is the result of `\"py\" in \"python\"`?", ["A) False", "B) True", "C) Index 0", "D) Error"], "B) True")
    ]

    # ---------- Interactive PowerPoint quiz ----------
    # The original deck showed the correct answer on every question slide.
    # This version hides the answer and makes A/B/C/D clickable in Slide Show mode.

    def add_box(slide, x, y, w, h, fill, line=None, radius=True):
        shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
        shape = slide.shapes.add_shape(shape_type, x, y, w, h)
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
        shape.line.color.rgb = line if line is not None else fill
        return shape

    def add_text(slide, text, x, y, w, h, size=20, bold=False,
                 color=TEXT_LIGHT, align=PP_ALIGN.CENTER):
        box = slide.shapes.add_textbox(x, y, w, h)
        tf = box.text_frame
        tf.word_wrap = True
        tf.clear()
        p = tf.paragraphs[0]
        p.text = text
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = color
        p.alignment = align
        return box

    def add_button(slide, text, x, y, w, h, target_slide,
                   fill=ACCENT_BLUE, font_size=20):
        button = add_box(slide, x, y, w, h, fill, fill)
        tf = button.text_frame
        tf.word_wrap = True
        tf.clear()
        p = tf.paragraphs[0]
        p.text = text
        p.font.size = Pt(font_size)
        p.font.bold = True
        p.font.color.rgb = TEXT_LIGHT
        p.alignment = PP_ALIGN.CENTER
        button.click_action.target_slide = target_slide
        return button

    # Create all slides first so internal links can point to already-created slides.
    title_slide = prs.slides.add_slide(blank_layout)
    question_slides = [prs.slides.add_slide(blank_layout) for _ in questions]
    correct_slides = [prs.slides.add_slide(blank_layout) for _ in questions]
    wrong_slides = [prs.slides.add_slide(blank_layout) for _ in questions]
    finish_slide = prs.slides.add_slide(blank_layout)

    # ----- Title / Start slide -----
    bg = add_box(title_slide, 0, 0, prs.slide_width, prs.slide_height, BG_DARK, BG_DARK, False)
    add_text(title_slide, "Python Basics - 100 MCQ Quiz",
             Inches(1), Inches(2.0), Inches(11.333), Inches(1.0),
             size=44, bold=True, color=ACCENT_BLUE)
    add_text(title_slide,
             "Variables, Strings, Conditionals, Lists, Tuples & Dictionaries",
             Inches(1), Inches(3.0), Inches(11.333), Inches(0.8),
             size=20, color=TEXT_LIGHT)
    add_button(title_slide, "START QUIZ  →",
               Inches(4.35), Inches(4.35), Inches(4.65), Inches(1.05),
               question_slides[0], fill=ACCENT_BLUE, font_size=24)
    add_text(title_slide, "Click Start Quiz, then choose A, B, C or D.",
             Inches(2.0), Inches(5.65), Inches(9.333), Inches(0.5),
             size=16, color=TEXT_LIGHT)

    # ----- Question slides -----
    positions = [
        (Inches(0.8), Inches(2.7)),
        (Inches(6.8), Inches(2.7)),
        (Inches(0.8), Inches(4.2)),
        (Inches(6.8), Inches(4.2))
    ]

    for idx, (q_text, opts, ans) in enumerate(questions):
        slide = question_slides[idx]

        add_box(slide, 0, 0, prs.slide_width, prs.slide_height,
                BG_DARK, BG_DARK, False)

        q_box = add_box(slide, Inches(0.8), Inches(0.6),
                        Inches(11.733), Inches(1.8), CARD_BG, ACCENT_BLUE)
        tf = q_box.text_frame
        tf.word_wrap = True
        tf.clear()
        p = tf.paragraphs[0]
        p.text = f"Q{idx + 1}. {q_text}"
        p.font.size = Pt(20)
        p.font.bold = True
        p.font.color.rgb = TEXT_LIGHT
        p.alignment = PP_ALIGN.CENTER

        # Every option links to either the correct or wrong feedback slide.
        for op_idx, opt in enumerate(opts):
            x, y = positions[op_idx]
            target = correct_slides[idx] if opt == ans else wrong_slides[idx]
            add_button(slide, opt, x, y, Inches(5.7), Inches(1.2),
                       target, fill=CARD_BG, font_size=18)

        add_text(slide, "Choose an answer",
                 Inches(4.5), Inches(6.05), Inches(4.333), Inches(0.45),
                 size=16, bold=True, color=TEXT_LIGHT)

    # ----- Feedback slides -----
    for idx, (q_text, opts, ans) in enumerate(questions):
        # Correct feedback
        slide = correct_slides[idx]
        add_box(slide, 0, 0, prs.slide_width, prs.slide_height,
                BG_DARK, BG_DARK, False)
        add_text(slide, "✓ CORRECT!", Inches(1), Inches(1.15),
                 Inches(11.333), Inches(1.0), size=44, bold=True,
                 color=ANSWER_BG)
        add_text(slide, f"Q{idx + 1}", Inches(1), Inches(2.2),
                 Inches(11.333), Inches(0.55), size=20, bold=True)
        add_text(slide, f"Correct Answer: {ans}",
                 Inches(1.5), Inches(3.0), Inches(10.333), Inches(1.0),
                 size=28, bold=True, color=TEXT_LIGHT)
        next_target = question_slides[idx + 1] if idx + 1 < len(questions) else finish_slide
        add_button(slide,
                   "NEXT QUESTION  →" if idx + 1 < len(questions) else "FINISH QUIZ  →",
                   Inches(4.0), Inches(4.8), Inches(5.333), Inches(1.0),
                   next_target, fill=ANSWER_BG, font_size=22)

        # Wrong feedback
        slide = wrong_slides[idx]
        add_box(slide, 0, 0, prs.slide_width, prs.slide_height,
                BG_DARK, BG_DARK, False)
        add_text(slide, "✗ WRONG ANSWER", Inches(1), Inches(1.15),
                 Inches(11.333), Inches(1.0), size=44, bold=True,
                 color=RGBColor(248, 113, 113))
        add_text(slide, f"Q{idx + 1}", Inches(1), Inches(2.2),
                 Inches(11.333), Inches(0.55), size=20, bold=True)
        add_text(slide, f"Correct Answer: {ans}",
                 Inches(1.5), Inches(3.0), Inches(10.333), Inches(1.0),
                 size=28, bold=True, color=TEXT_LIGHT)
        add_button(slide,
                   "NEXT QUESTION  →" if idx + 1 < len(questions) else "FINISH QUIZ  →",
                   Inches(4.0), Inches(4.8), Inches(5.333), Inches(1.0),
                   next_target, fill=ACCENT_BLUE, font_size=22)

    # ----- Finish slide -----
    add_box(finish_slide, 0, 0, prs.slide_width, prs.slide_height,
            BG_DARK, BG_DARK, False)
    add_text(finish_slide, "🎉 QUIZ COMPLETE!",
             Inches(1), Inches(2.0), Inches(11.333), Inches(1.0),
             size=44, bold=True, color=ACCENT_BLUE)
    add_text(finish_slide,
             "You have reached the end of the 100-question Python quiz.",
             Inches(1.5), Inches(3.15), Inches(10.333), Inches(0.8),
             size=22, color=TEXT_LIGHT)
    add_button(finish_slide, "RESTART QUIZ  ↻",
               Inches(4.35), Inches(4.6), Inches(4.65), Inches(1.0),
               question_slides[0], fill=ACCENT_BLUE, font_size=22)

    prs.save("Python_100_MCQ_Quiz_Interactive.pptx")
    print("Interactive PowerPoint generated successfully: Python_100_MCQ_Quiz_Interactive.pptx")

create_python_quiz_ppt()